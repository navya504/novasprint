from pathlib import Path
import shutil, textwrap
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
DOCS = ROOT / 'docs'
DOCS.mkdir(exist_ok=True)
OUT = DOCS / 'NovaSprint_Implementation_Guide.docx'
TMP = ROOT / '.doc-build'
TMP.mkdir(exist_ok=True)
NAVY='17365F'; BLUE='2F6B9C'; TEAL='29A6A2'; PURPLE='7464C8'; CORAL='E88067'; GOLD='D5A342'; INK='26364A'; MUTED='61758D'; PALE='F3F7FC'; LINE='DCE5EF'; WHITE='FFFFFF'

def font_path(bold=False):
    p=Path(r'C:\Windows\Fonts\segoeuib.ttf' if bold else r'C:\Windows\Fonts\segoeui.ttf')
    return str(p) if p.exists() else None

def get_font(size,bold=False):
    path=font_path(bold)
    return ImageFont.truetype(path,size) if path else ImageFont.load_default()

def wrap_draw(draw,text,font,x,y,width,fill,leading=1.3,max_lines=None):
    words=str(text).split(); lines=[]; line=''
    for word in words:
        trial=(line+' '+word).strip()
        if draw.textbbox((0,0),trial,font=font)[2] <= width or not line: line=trial
        else: lines.append(line); line=word
    if line: lines.append(line)
    if max_lines and len(lines)>max_lines:
        lines=lines[:max_lines]; lines[-1]=lines[-1].rstrip(' .')+'…'
    h=draw.textbbox((0,0),'Ag',font=font)[3]-draw.textbbox((0,0),'Ag',font=font)[1]
    for s in lines: draw.text((x,y),s,font=font,fill=fill); y+=int(h*leading)
    return y

def arrow(draw,a,b,color=BLUE,width=7):
    draw.line([a,b],fill='#'+color,width=width)
    import math
    ang=math.atan2(b[1]-a[1],b[0]-a[0]); length=18; spread=.56
    p1=(b[0]-length*math.cos(ang-spread),b[1]-length*math.sin(ang-spread))
    p2=(b[0]-length*math.cos(ang+spread),b[1]-length*math.sin(ang+spread))
    draw.polygon([b,p1,p2],fill='#'+color)

def phase_bar(draw,y,label,detail):
    draw.rounded_rectangle((55,y,1545,y+44),radius=18,fill='#EAF0F8')
    draw.text((78,y+10),label,font=get_font(19,True),fill='#'+NAVY)
    draw.text((420,y+12),detail,font=get_font(16),fill='#'+MUTED)

def node(draw,x,y,w,h,num,title,detail,color):
    draw.rounded_rectangle((x+5,y+7,x+w+5,y+h+7),radius=22,fill='#E2E9F2')
    draw.rounded_rectangle((x,y,x+w,y+h),radius=22,fill='#'+WHITE,outline='#D8E2ED',width=2)
    draw.rounded_rectangle((x+18,y+18,x+68,y+68),radius=15,fill='#'+color)
    draw.text((x+32,y+29),str(num),font=get_font(22,True),fill='#'+WHITE)
    y2=wrap_draw(draw,title,get_font(21,True),x+82,y+20,w-100,'#'+NAVY,1.15,2)
    wrap_draw(draw,detail,get_font(15),x+22,max(y+88,y2+7),w-44,'#'+MUTED,1.28,3)

def make_workflow(path):
    im=Image.new('RGB',(1600,1000),'#F9FBFE'); d=ImageDraw.Draw(im)
    d.text((58,37),'NovaSprint planning graph',font=get_font(31,True),fill='#'+NAVY)
    d.text((60,79),'A single, evidence-led pass from sprint goal to Jira delivery',font=get_font(18),fill='#'+MUTED)
    phase_bar(d,118,'01  CONTEXT & SCOPE','Bring the goal, people, work history, and priorities into a shared plan state')
    phase_bar(d,385,'02  PLAN & ASSESS','Shape delivery, forecast effort, allocate owners, monitor evidence, assess outcomes')
    phase_bar(d,652,'03  SYNTHESIZE & DELIVER','Summarize the decision and publish planned work to Jira')
    w=325; h=166; xs=[60,445,830,1215]; y1=174; y2=441; y3=708
    colors=[BLUE,TEAL,PURPLE,BLUE,GOLD,TEAL,BLUE,PURPLE,CORAL,NAVY]
    top=[('Team Context Agent','Reads confirmed team profiles, Jira issues, board and sprint context.'),('Intake Agent','Normalizes the goal, dates, backlog, priorities and constraints.'),('Task Decomposer Agent','Derives goal-driven tasks, dependencies and acceptance criteria.'),('Sprint Planner Agent','Builds candidate scopes, selects one and orders delivery.')]
    for i,(title,detail) in enumerate(top): node(d,xs[i],y1,w,h,i+1,title,detail,colors[i])
    for i in range(3): arrow(d,(xs[i]+w+8,y1+h//2),(xs[i+1]-12,y1+h//2))
    arrow(d,(xs[3]+w//2,y1+h+7),(xs[3]+w//2,y2-12))
    bottom=[('Predictor Agent','Forecasts time, resources, outcome likelihood and rework risk.'),('Assignment Agent','Matches work to skills, availability, preferences and workload.'),('Monitor Agent','Reports observed Jira progress and schedule deviations.'),('Evaluator Agent','Assesses goal fit and evaluates supplied user-impact evidence.')]
    # Second row is traversed right to left: step 5 starts at the right edge.
    for j,(title,detail) in enumerate(bottom):
        pos=3-j; node(d,xs[pos],y2,w,h,5+j,title,detail,colors[4+j])
    for pos in [3,2,1]: arrow(d,(xs[pos]-10,y2+h//2),(xs[pos-1]+w+12,y2+h//2))
    arrow(d,(xs[0]+w//2,y2+h+7),(xs[0]+w//2+210,y3-13))
    node(d,375,y3,360,h,9,'Orchestrator Agent','Combines the selected roadmap, forecasts, owners, QA, release and evaluation outputs.',colors[8])
    node(d,865,y3,455,h,10,'Jira Publisher Agent','Creates tasks and subtasks, moves issues to To Do, assigns mapped owners, and links planned work.',colors[9])
    arrow(d,(740,y3+h//2),(850,y3+h//2))
    d.text((60,914),'The graph has explicit sequential edges. Intake validation can stop a run early; an empty decomposition gets one focused retry.',font=get_font(15),fill='#'+MUTED)
    im.save(path)

def make_paper_map(path):
    """Illustrate the six-module MAP algorithm described in Webb et al. (2025)."""
    im=Image.new('RGB',(1600,1000),'#F8FAFD'); d=ImageDraw.Draw(im)
    d.text((55,32),'Modular Agentic Planner (MAP)',font=get_font(32,True),fill='#17365F')
    d.text((58,78),'Paper architecture: subgoals guide an action loop, state prediction and evaluation, and goal coordination',font=get_font(17),fill='#61758D')
    # The search area groups state prediction, value estimation and action selection.
    d.rounded_rectangle((690,330,1545,890),radius=26,fill='#F9E7E3',outline='#E88067',width=3)
    d.text((715,350),'SEARCH',font=get_font(18,True),fill='#C95F4B')
    d.rounded_rectangle((315,535,910,790),radius=22,fill='#DDF0F2',outline='#42A9B0',width=3)
    d.text((340,553),'ACTION PROPOSAL LOOP',font=get_font(16,True),fill='#277F87')
    # Goal and starting state.
    d.rounded_rectangle((35,395,235,565),radius=22,fill='#FFFFFF',outline='#D8E2ED',width=2)
    d.text((58,418),'Goal',font=get_font(21,True),fill='#17365F')
    d.text((58,450),'s_goal',font=get_font(18),fill='#2F6B9C')
    d.text((58,490),'Initial state',font=get_font(17,True),fill='#17365F')
    d.text((58,520),'s₀',font=get_font(18),fill='#2F6B9C')
    node(d,315,145,300,145,1,'Task Decomposer','Turns the goal and initial state into subgoals.','7464C8')
    node(d,1245,145,250,145,6,'Orchestrator','Checks whether a subgoal or the final goal is complete.','17365F')
    d.rounded_rectangle((1520,178,1585,250),radius=17,fill='#FFFFFF',outline='#D8E2ED',width=2)
    d.text((1530,201),'Plan',font=get_font(16,True),fill='#17365F')
    # Action proposal modules.
    node(d,370,620,220,145,2,'Actor','Proposes actions for the current state and subgoal.','2F6B9C')
    node(d,650,620,220,145,3,'Monitor','Checks action validity and returns constraint feedback.','29A6A2')
    # Search modules and action choice.
    node(d,775,430,245,145,4,'Predictor','Predicts the next state after a proposed action.','E88067')
    node(d,1135,430,245,145,5,'Evaluator','Estimates the predicted state’s value toward the goal.','D5A342')
    d.rounded_rectangle((1110,675,1405,815),radius=20,fill='#FFFFFF',outline='#D8E2ED',width=2)
    d.rounded_rectangle((1130,694,1180,744),radius=14,fill='#7464C8')
    d.text((1144,705),'S',font=get_font(20,True),fill='#FFFFFF')
    d.text((1195,697),'Selection',font=get_font(19,True),fill='#17365F')
    wrap_draw(d,'Chooses the action that best advances the active subgoal.',get_font(15),1130,754,255,'#61758D',1.2,2)
    # Inputs and subgoal handoff.
    arrow(d,(235,445),(300,265),'2F6B9C',6)
    arrow(d,(615,215),(1230,215),'7464C8',6)
    d.text((850,184),'subgoals',font=get_font(15,True),fill='#7464C8')
    arrow(d,(1499,215),(1515,215),'17365F',6)
    # Orchestrator supplies an active subgoal to the search and receives the selected result.
    d.line([(1270,290),(1270,365),(735,365),(735,600)],fill='#7464C8',width=5)
    d.polygon([(735,600),(725,582),(745,582)],fill='#7464C8')
    d.text((900,338),'active subgoal',font=get_font(14,True),fill='#7464C8')
    d.line([(1405,745),(1480,745),(1480,300)],fill='#17365F',width=6)
    d.polygon([(1480,290),(1470,308),(1490,308)],fill='#17365F')
    d.text((1410,680),'selected result',font=get_font(13,True),fill='#17365F')
    # Actor and Monitor feedback loop.
    arrow(d,(590,690),(640,690),'2F6B9C',7)
    d.text((555,660),'candidate actions',font=get_font(13,True),fill='#2F6B9C')
    d.line([(650,755),(620,825),(480,825),(480,780)],fill='#29A6A2',width=5)
    d.polygon([(480,780),(470,798),(490,798)],fill='#29A6A2')
    d.text((500,798),'invalid action feedback',font=get_font(13,True),fill='#277F87')
    # Candidate action to next-state prediction, then value-guided selection.
    d.line([(870,690),(940,690),(940,582)],fill='#E88067',width=6)
    d.polygon([(940,582),(930,600),(950,600)],fill='#E88067')
    d.text((960,650),'valid action + state',font=get_font(13,True),fill='#C95F4B')
    arrow(d,(1020,500),(1125,500),'E88067',7)
    d.text((1025,468),'predicted state',font=get_font(13,True),fill='#C95F4B')
    d.line([(1260,575),(1260,645)],fill='#D5A342',width=6)
    d.polygon([(1260,645),(1250,627),(1270,627)],fill='#D5A342')
    d.text((1275,603),'state value',font=get_font(13,True),fill='#A77C20')
    # Orchestration sends the active subgoal back into the proposal loop until complete.
    d.text((58,940),'MAP paper modules and control flow. NovaSprint’s production workflow is related but does not implement these action and search loops.',font=get_font(14),fill='#61758D')
    im.save(path)

def card(draw,box,title,body,color,subtitle=None):
    x,y,w,h=box
    draw.rounded_rectangle((x+4,y+7,x+w+4,y+h+7),radius=24,fill='#E2E9F2')
    draw.rounded_rectangle((x,y,x+w,y+h),radius=24,fill='#'+WHITE,outline='#D8E2ED',width=2)
    draw.rounded_rectangle((x,y,x+w,y+16),radius=8,fill='#'+color)
    yy=wrap_draw(draw,title,get_font(22,True),x+24,y+34,w-48,'#'+NAVY,1.1,2)
    if subtitle: yy=wrap_draw(draw,subtitle,get_font(15,True),x+24,yy+7,w-48,'#'+color,1.15,2)
    wrap_draw(draw,body,get_font(16),x+24,yy+10,w-48,'#'+MUTED,1.27,4)

def make_architecture(path):
    im=Image.new('RGB',(1600,900),'#F9FBFE'); d=ImageDraw.Draw(im)
    d.text((58,36),'NovaSprint runtime architecture',font=get_font(31,True),fill='#'+NAVY)
    d.text((60,80),'The browser sends requests to a local service that coordinates agents and server-side integrations',font=get_font(17),fill='#'+MUTED)
    card(d,(60,185,350,205),'Browser UI','HTML, CSS and JavaScript. Collects sprint inputs, streams agent events, renders plans and tracking dashboards.',BLUE,'User interaction')
    card(d,(625,150,390,235),'Node.js HTTP API','Serves the app, validates requests, manages team/settings/history, streams planning runs and stores completed plans.',TEAL,'Local service · 127.0.0.1:4174 by default')
    card(d,(1190,185,350,205),'LangGraph planning graph','Typed shared state and explicit graph edges coordinate NovaSprint roles. Zod schemas validate structured outputs.',PURPLE,'Sequential workflow')
    arrow(d,(420,286),(610,286),BLUE,8); d.text((454,255),'HTTP + SSE',font=get_font(15,True),fill='#'+BLUE)
    arrow(d,(1020,286),(1175,286),TEAL,8); d.text((1035,255),'invoke graph',font=get_font(15,True),fill='#'+TEAL)
    # These cards show server-side dependencies. Keeping them grouped avoids implying
    # a separate service or a particular internal call path that the code does not have.
    d.rounded_rectangle((38,460,1562,843),radius=30,fill='#F1F5FA',outline='#D7E1ED',width=2)
    d.text((60,478),'SERVER-SIDE DEPENDENCIES',font=get_font(17,True),fill='#'+NAVY)
    d.text((60,506),'The API and planning graph use these configured services and local data.',font=get_font(15),fill='#'+MUTED)
    card(d,(60,600,425,205),'Nebius or Ollama','Nebius Token Factory uses the OpenAI-compatible LangChain adapter. Local Ollama is an alternative provider. Credentials stay server-side.',GOLD,'Model gateway')
    card(d,(590,600,425,205),'Jira Cloud REST + Agile APIs','Reads project, board and sprint context; creates issues and subtasks; assigns users, transitions to To Do, adds sprint items and links dependencies.',CORAL,'Project system of record')
    card(d,(1120,600,420,205),'JSON workspace store','seed.json initializes the workspace. workspace.json persists settings, team profiles, plans and audit events with an atomic temp-file replacement.',BLUE,'Local persistence')
    d.rounded_rectangle((35,130,1565,843),radius=32,outline='#D7E1ED',width=2)
    # Redraw foreground cards and arrows over the boundary frame.
    card(d,(60,185,350,205),'Browser UI','HTML, CSS and JavaScript. Collects sprint inputs, streams agent events, renders plans and tracking dashboards.',BLUE,'User interaction')
    card(d,(625,150,390,235),'Node.js HTTP API','Serves the app, validates requests, manages team/settings/history, streams planning runs and stores completed plans.',TEAL,'Local service · 127.0.0.1:4174 by default')
    card(d,(1190,185,350,205),'LangGraph planning graph','Typed shared state and explicit graph edges coordinate NovaSprint roles. Zod schemas validate structured outputs.',PURPLE,'Sequential workflow')
    arrow(d,(420,286),(610,286),BLUE,8); d.text((454,255),'HTTP + SSE',font=get_font(15,True),fill='#'+BLUE)
    arrow(d,(1020,286),(1175,286),TEAL,8); d.text((1035,255),'invoke graph',font=get_font(15,True),fill='#'+TEAL)
    im.save(path)

def make_tracking(path):
    im=Image.new('RGB',(1600,600),'#F9FBFE'); d=ImageDraw.Draw(im)
    d.text((58,34),'Tracking is a separate progress workflow',font=get_font(30,True),fill='#'+NAVY)
    d.text((60,77),'It checks Jira state for a saved plan; it does not rerun sprint decomposition or planning.',font=get_font(17),fill='#'+MUTED)
    boxes=[(50,'Plan History','User selects Track for one saved plan.',BLUE),(365,'Jira lookup','Server fetches issues carrying the plan label.',CORAL),(680,'Monitor Agent','Summarizes observed issue status, KPIs and deviations.',TEAL),(995,'Orchestrator summary','A compact tracking summary is also generated and stored.',PURPLE),(1310,'Dashboard','UI shows the Monitor output, completion bar and Jira task evidence.',NAVY)]
    for n,(x,title,body,color) in enumerate(boxes,1):
        card(d,(x,190,245,240),title,body,color,f'Step {n}')
    for i in range(4): arrow(d,(boxes[i][0]+251,310),(boxes[i+1][0]-12,310),[BLUE,CORAL,TEAL,PURPLE][i],7)
    d.rounded_rectangle((58,478,1540,555),radius=18,fill='#EAF0F8')
    d.text((83,496),'Completion counts use Jira’s Done category or a resolution date. KPI comparisons appear only when metric configuration and evidence exist.',font=get_font(16,True),fill='#'+NAVY)
    im.save(path)

def shade(cell,fill):
    pr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); pr.append(shd)

def set_cell_margins(cell,top=100,start=120,bottom=100,end=120):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
    for edge,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=OxmlElement('w:'+edge); node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa'); mar.append(node)
    tcPr.append(mar)

def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def set_cell_text(cell,text,bold=False,color=INK,size=9.2):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(0)
    r=p.add_run(str(text)); r.bold=bold; r.font.name='Aptos'; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cell)

def add_table(doc,headers,rows,widths=None,font_size=9.1):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    if widths:
        for i,w in enumerate(widths): t.columns[i].width=Inches(w)
    for i,h in enumerate(headers): shade(t.rows[0].cells[i],NAVY); set_cell_text(t.rows[0].cells[i],h,True,WHITE,font_size)
    set_repeat_table_header(t.rows[0])
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            if widths: cells[i].width=Inches(widths[i])
            shade(cells[i],WHITE if ridx%2 else PALE); set_cell_text(cells[i],val,i==0,NAVY if i==0 else INK,font_size)
        trPr=t.rows[-1]._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

def add_callout(doc,title,text,color=TEAL):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(8); p.paragraph_format.line_spacing=1.15
    r=p.add_run(f'{title}. '); r.bold=True; r.font.name='Aptos'; r.font.size=Pt(10.5); r.font.color.rgb=RGBColor.from_string(INK)
    r=p.add_run(text); r.font.name='Aptos'; r.font.size=Pt(10.5); r.font.color.rgb=RGBColor.from_string(INK)

def add_p(doc,text,style=None):
    p=doc.add_paragraph(style=style); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.12; p.add_run(text); return p

def add_heading(doc,text,level=1):
    p=doc.add_heading(text,level=level); p.paragraph_format.keep_with_next=True; return p

def add_figure(doc,path,caption,width=6.85):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(3)
    r=p.add_run(); r.add_picture(str(path),width=Inches(width))
    cap=doc.add_paragraph(); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap.paragraph_format.space_after=Pt(8)
    rr=cap.add_run(caption); rr.italic=True; rr.font.size=Pt(8.3); rr.font.color.rgb=RGBColor.from_string(MUTED)

def add_bullets(doc,items):
    for item in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.first_line_indent=Inches(-.12); p.add_run(item)

def page_number(paragraph):
    paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=paragraph.add_run('NovaSprint  •  Implementation Guide  |  '); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); paragraph._p.append(fld)

paper_map=TMP/'paper_map_architecture.png'; workflow=TMP/'novasprint_workflow.png'; architecture=TMP/'system_architecture.png'; tracking=TMP/'tracking_flow.png'
make_paper_map(paper_map); make_workflow(workflow); make_architecture(architecture); make_tracking(tracking)

doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.64); sec.bottom_margin=Inches(.66); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
styles=doc.styles
normal=styles['Normal']; normal.font.name='Aptos'; normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string(INK); normal.paragraph_format.space_after=Pt(6)
for style_name,size,color in [('Title',30,'000000'),('Heading 1',20,NAVY),('Heading 2',13,BLUE),('Heading 3',10,TEAL)]:
    st=styles[style_name]; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.space_before=Pt(8); st.paragraph_format.space_after=Pt(5)
styles['List Bullet'].font.name='Aptos'; styles['List Bullet'].font.size=Pt(9.2)
page_number(sec.footer.paragraphs[0])

# Cover and opening rationale
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(4)
r=p.add_run('NOVASPRINT  /  TECHNICAL GUIDE'); r.bold=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(TEAL); r.font.letter_spacing=Pt(1) if hasattr(r.font,'letter_spacing') else None
p=doc.add_paragraph(style='Title'); p.paragraph_format.space_after=Pt(3); p.add_run('NovaSprint Implementation Guide')
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(14); rr=p.add_run('Brain-inspired MAP planning architecture and technical implementation'); rr.font.size=Pt(13); rr.font.color.rgb=RGBColor.from_string(MUTED)
add_heading(doc,'Why sprint planning matters',1)
add_p(doc,'Planning matters because a goal alone does not tell a team what to build first, how work depends on other work, who can own it, how much time is plausible, or how success will be checked. Without those decisions, teams can ship disconnected work, discover integration or QA gaps late, or call a release complete without measuring the user outcome.')
add_p(doc,'NovaSprint turns a sprint goal into an evidence-informed delivery plan. It combines team skills and availability with Jira project context, then separates task design, forecasting, ownership, progress monitoring, outcome evaluation, and final synthesis into distinct agent roles. The plan flows to Jira as work items and can be tracked against live Jira status.')
add_callout(doc,'The operating idea','Convert intent into an ordered, owned, measurable plan, while keeping forecasts distinct from observed results.',TEAL)
add_heading(doc,'What the implementation provides',2)
add_table(doc,['Planning need','NovaSprint response'],[
 ('A goal that is too broad to execute','Task Decomposer derives a Jira-ready task list from the goal, even when no backlog list is supplied.'),
 ('Work that can be sequenced and owned','Sprint Planner selects a goal-aligned scope; Predictor estimates; Assignment matches roles, skills and availability.'),
 ('Quality and release coverage','The decomposition includes one goal-level end-to-end QA task and one sprint release-readiness task when applicable.'),
 ('Progress that can be checked later','Track queries the plan’s Jira issues and shows Monitor output, task evidence and completion status.')
],widths=[2.03,4.72],font_size=8.9)
add_p(doc,'This guide describes the implementation represented by the NovaSprint source code as of October 2026. It separates the research paper’s Modular Agentic Planner from NovaSprint’s related, sequential LangGraph workflow so the architecture and its limits are clear.')

# Agent primer and framework requested for product implementation review.
doc.add_page_break(); add_heading(doc,'The agent primer',1)
add_p(doc,'Before choosing an architecture, state what the agent does autonomously, what it hands to a person, and how the end-to-end workflow will be measured. The one-line definition below describes NovaSprint as it is implemented today; unknown targets are left explicit instead of guessed.')
add_heading(doc,'NovaSprint in one line',2)
add_p(doc,'NovaSprint helps product teams turn a one-line sprint goal into an estimated, assigned and trackable Jira plan in a local web app, replacing manual task breakdown, forecasting, assignment and Jira entry; it reads team and Jira context, uses two integration categories (a configured LLM provider and Jira Cloud), publishes planned Jira work automatically, and hands the result to the team to inspect and track; end-to-end completion time and success-rate targets have not yet been measured.')
add_heading(doc,'One-line template',2)
add_p(doc,'My agent helps [USER] do [MULTI-STEP TASK] in [SURFACE], replacing [MANUAL WORKFLOW AND COST]. It does [WORK] on its own using [N TOOLS], hands off to a human [WHEN], and I’ll know it works when [USER] can [COMPLETE THE TASK] in under [TIME] with [SUCCESS RATE].')
add_heading(doc,'Worked example',2)
add_p(doc,'My agent helps a startup founder research a new market in a web app, replacing the 6 hours of manual googling and tab-juggling it takes to size up competitors today. It searches the web, pulls each competitor’s pricing and positioning, and drafts a one-page brief on its own using 4 tools, hands off to the founder to review before the brief gets saved, and I’ll know it works when a founder can get a usable market brief in under 15 minutes that they would actually send to their team 8 times out of 10.')
add_heading(doc,'Three rules for the one-liner',2)
add_table(doc,['Rule','Meaning'],[
 ('Measure task completion, not one-call accuracy.','The agent succeeds when a user finishes the full workflow. Measure the end-to-end result.'),
 ('Make state explicit.','Specify what the agent remembers, for how long, and where it lives. Session state, conversation history and persistent storage are different.'),
 ('Require human review for writes by default.','Creating, modifying, sending or paying should normally pause for a human. Reads can run autonomously; write access needs an explicit policy.')
],widths=[2.0,4.75],font_size=8.8)
add_p(doc,'NovaSprint product-policy note. The current implementation automatically publishes the generated plan to Jira when planning finishes, as requested for this product. It therefore has no pre-publication human approval step and is an explicit exception to the default write-review rule above.')

doc.add_page_break(); add_heading(doc,'NovaSprint agent framework',1)
add_p(doc,'The framework below distinguishes the current implementation from measurement targets that still need to be set. It is grounded in the source code and does not claim unmeasured time savings or success rates.')
add_table(doc,['Field','NovaSprint definition'],[
 ('Agent goal','Convert a sprint goal and project context into an ordered, estimated, assigned plan; create the planned Jira issues; then let the team track Jira progress.'),
 ('Where people use it','A local browser app served by the Node.js service. Planning output, plan history, Jira publication and Track progress are available in the web interface.'),
 ('Steps, in order','Read team/Jira context → normalize the request → decompose the goal into tasks → choose and order scope → predict effort/outcomes → assign owners → assess observed Jira context and goal impact → synthesize the plan → publish Jira issues. Track is a separate on-demand monitor workflow.'),
 ('What it can do','Read Jira project, board, sprint and issue context; read local team and workspace state; call the configured Nebius or Ollama model; write plan/history state; create Jira issues and subtasks, assign mapped accounts, transition to To Do, add sprint membership and link dependencies.'),
 ('What it remembers','During a run, LangGraph state holds the request, Jira snapshot, work items, predictions, assignments and agent outputs. Across runs, data/workspace.json persists settings, team profiles, saved plans, tracking reports and audit events; data/seed.json initializes a workspace.'),
 ('What it must never do','Invent task results, baseline or post-change metrics, outcome percentages, Jira statuses or evidence. It should not silently substitute a fixed sample plan or hard-coded response when agents fail; failures should be explicit.'),
 ('Human involvement','Current product policy publishes Jira work automatically after planning, then gives the team the plan to inspect and Track. There is no pre-write approval gate in this build; the default write-safety rule recommends adding one if policy changes.'),
 ('When something breaks','Structured model output gets one compact-context retry with a larger response budget; empty decomposition gets one focused retry. A second failure is surfaced as an error. Jira publication reports partial outcomes when individual writes fail; the planning graph does not loop from evaluation back to task generation.'),
 ('How success is measured','Primary proposed measure: end-to-end plan-and-publish success rate—the share of valid goal runs that produce a complete task list and create every planned Jira issue that can then be checked by Track. Planning-time target and success-rate threshold are not yet instrumented or supplied.')
],widths=[1.72,5.03],font_size=8.25)
add_p(doc,'Measurement status: the code provides task counts, Jira publication outcomes, elapsed agent durations and tracking results, but it does not yet calculate a product-level completion-time baseline or an accepted-plan success rate. Set those thresholds after collecting real runs.')

# MAP flow
doc.add_page_break(); add_heading(doc,'1  MAP planning architecture',1)
add_p(doc,'MAP means Modular Agentic Planner in the paper by Webb, Mondal and Momennejad. It decomposes a goal into subgoals, proposes and monitors candidate actions, predicts resulting states, evaluates those states, selects actions through search, and uses an Orchestrator to decide when subgoals and the final goal are reached. The colored modules below reproduce those responsibilities and feedback relationships at a product-guide level.')
add_figure(doc,paper_map,'Figure 1. The research paper’s MAP modules: decomposition, monitored action proposals, state prediction and evaluation, selection, and orchestration.',6.88)
add_p(doc,'Source: Taylor Webb, Shanka Subhra Mondal and Ida Momennejad, “A brain-inspired agentic architecture to improve planning with LLMs,” Nature Communications 16, 8633 (2025), https://doi.org/10.1038/s41467-025-63804-5. Microsoft Research publication page: https://www.microsoft.com/en-us/research/publication/a-brain-inspired-agentic-architecture-to-improve-planning-with-llms/. The MAP architecture and reported evaluations belong to the cited paper; they are not NovaSprint benchmark results.')
add_heading(doc,'How NovaSprint relates to the paper',2)
add_p(doc,'NovaSprint applies the idea of specialist roles coordinated around a goal, but its current graph is not the paper’s full MAP algorithm. It is a mostly sequential LangGraph workflow with one explicit task-decomposition retry; it does not implement the paper’s repeated Actor–Monitor action loop or Predictor–Evaluator tree search. The differences below matter because similarly named roles have different jobs.')
add_table(doc,['Paper MAP module','NovaSprint relationship'],[
 ('Task Decomposer','Closest conceptual match: derives goal-based work items, dependencies and acceptance criteria.'),
 ('Actor','No separate Actor module in the current planning graph. Sprint Planner selects and orders task scope; Jira Publisher performs configured Jira writes.'),
 ('Monitor','Different responsibility. Paper MAP Monitor gates proposed actions against task constraints and feeds back errors; NovaSprint Monitor Agent reports Jira task progress and schedule/KPI deviations.'),
 ('Predictor','Different output. Paper MAP predicts the next state after an action; NovaSprint Predictor forecasts task time, resources, outcome likelihood and rework risk.'),
 ('Evaluator','Different output. Paper MAP values a predicted state toward a goal; NovaSprint Evaluator assesses usability, goal alignment and supplied outcome evidence. It does not rank plans.'),
 ('Orchestrator','Paper MAP decides whether a subgoal/final goal is achieved during iterative planning; NovaSprint Orchestrator synthesizes the completed plan and the separate tracking workflow summary.')
],widths=[1.5,5.25],font_size=8.3)
add_heading(doc,'NovaSprint planning graph',2)
add_figure(doc,workflow,'Figure 2. NovaSprint’s implemented planning flow: eight model-backed roles plus deterministic context collection and Jira publishing stages.',6.88)
add_p(doc,'The paper’s architecture is the conceptual reference. LangGraph is the execution framework used by this application; the current LangGraph graph follows NovaSprint’s product flow and should not be described as a faithful implementation of MAP’s internal action proposal and search loops.')
add_heading(doc,'Execution characteristics',2)
add_bullets(doc,[
 'The main planning path is sequential and has explicit edges. It does not cycle from evaluation back to task generation.',
 'Invalid sprint goals or date ranges can stop the graph at intake. An empty task decomposition gets one focused model retry, then returns an error rather than silently inventing tasks.',
 'Each agent receives a role-specific prompt and structured output schema. Length-limit errors trigger one compact-context retry with a larger response budget.',
 'The final Orchestrator output summarizes the selected roadmap and agent findings. Jira Publisher runs as the graph’s final delivery stage.'
])

# Agent roles, part 1
doc.add_page_break(); add_heading(doc,'2  Agent responsibilities',1)
add_p(doc,'Each agent has a bounded responsibility. That separation keeps the forecast, observed progress, impact assessment and plan selection from being blended into one ambiguous summary.')
agents_a=[
 ('Team Context Agent','Reads the selected Jira project, board, sprint list and team profiles.','Jira snapshot, teammate skills and availability, existing work evidence.'),
 ('Intake Agent','Normalizes the sprint goal, dates, priorities, constraints and optional backlog.','A validated brief that preserves the request and its boundaries.'),
 ('Task Decomposer Agent','Turns the goal into concrete delivery tasks, dependencies and acceptance criteria.','Non-empty, Jira-ready task tree; goal-level QA and release tasks are distinct items.'),
 ('Sprint Planner Agent','Builds up to two scope options, recommends one and orders selected work.','Selected task IDs, sequence, rationale and tradeoffs; release readiness is sequenced last.'),
 ('Predictor Agent','Forecasts task time, resources, outcome probability and rework risk using available history and complexity.','Per-task forecasts, confidence and explained evidence limitations; estimates are not commitments.')]
add_table(doc,['Agent','Responsibility','Output'],agents_a,widths=[1.45,2.75,2.55],font_size=8.2)
add_heading(doc,'Planning contract',2)
add_p(doc,'The Decomposer owns task creation. The Sprint Planner owns scope selection. Predictor owns forecasts. Assignment owns proposed owners. Evaluator assesses user impact and measurement evidence; it does not rank plan options. Monitor reports observed task progress and deviations; it does not predict completion time.')
add_callout(doc,'Guardrail','The Evaluator reports baseline/post-change comparisons only when both values and their source are provided. Before release, impact can be described as an evaluation plan, but it is not reported as a measured result.',BLUE)

# Agent roles, part 2
doc.add_page_break(); add_heading(doc,'2  Agent responsibilities continued',1)
agents_b=[
 ('Assignment Agent','Matches task needs to configured team roles, confirmed skills, availability, preferences and workload.','Named owners with rationale, or an explicit unassigned result when evidence is not sufficient.'),
 ('Monitor Agent','Reports observed Jira status, task evidence, configured KPI movement and deviations.','A concise status report and task-level Jira evidence. Tracking uses a separate progress flow.'),
 ('Evaluator Agent','Assesses the goal’s expected usability and user impact and how to measure it.','Goal-alignment and usability findings; metric comparison only from supplied baseline and post-change evidence.'),
 ('Orchestrator Agent','Integrates the chosen roadmap, Predictor estimates, Assignment owners, QA, release, evaluation and monitoring summaries.','A concise final planning blueprint using task and teammate names.'),
 ('Jira Publisher Agent','Creates plan tasks and subtasks, assigns mapped accounts, sets To Do, adds configured sprint membership and links dependencies.','Jira issue keys and operation results; partial publication is reported if an operation fails.')]
add_table(doc,['Agent','Responsibility','Output'],agents_b,widths=[1.45,2.75,2.55],font_size=8.2)
add_heading(doc,'Quality tasks are in the plan, not scattered everywhere',2)
add_p(doc,'For applicable goals, the Decomposer creates one comprehensive end-to-end QA task and one goal-level release-readiness plan. This keeps validation and release ownership explicit without duplicating QA and release subtasks under every implementation task. UI redesign goals also call for frontend delivery and backend compatibility work.')
add_heading(doc,'Estimate versus observation',2)
add_table(doc,['Signal','Owner','Evidence rule'],[
 ('Time, resources, outcome probability, rework risk','Predictor Agent','Forecasts are marked with confidence and grounded in available velocity, capacity, Jira history and task complexity.'),
 ('Current Jira status, completion, deviations','Monitor Agent','Observed from Jira issue fields and timestamps during a Track run.'),
 ('Usability or KPI outcome after the change','Evaluator Agent','Requires supplied baseline and post-change evidence; no fabricated before/after values.')
],widths=[2.18,1.38,3.19],font_size=8.5)

# System architecture
doc.add_page_break(); add_heading(doc,'3  System architecture',1)
add_p(doc,'The browser owns presentation and user input. A local Node.js server owns API routes, configuration, model and Jira credentials, graph execution and persistence. The browser never calls the model provider or Jira with server credentials.')
add_figure(doc,architecture,'Figure 3. NovaSprint’s local service coordinates the graph and keeps model, Jira and workspace access on the server side.',6.88)
add_table(doc,['Layer','Implementation','Key source files'],[
 ('Web interface','Single-page HTML/CSS with browser JavaScript for team, history, planner, agent output, export and tracking views.','index.html · styles.css · app.js · planner-ui.js'),
 ('HTTP service','Node.js HTTP server serves static assets and JSON/SSE endpoints; binds to 127.0.0.1 by default.','server.js'),
 ('Agent orchestration','LangGraph StateGraph with shared annotations, explicit edges, streamed run events and conditional early exit.','lib/graph.js'),
 ('Structured contracts','Zod schemas validate agent outputs for intake, decomposition, planning, prediction, assignment, monitoring, evaluation and orchestration.','lib/agent-schemas.js'),
 ('Model adapters','Nebius Token Factory through LangChain ChatOpenAI, or local Ollama through ChatOllama.','lib/model.js'),
 ('Jira adapter','Jira Cloud REST and Agile APIs for project/board/sprint context, issue creation, assignment, status transitions and links.','lib/jira.js'),
 ('Persistence','Seeded JSON workspace store with atomic file replacement and bounded audit history.','lib/store.js · data/seed.json')
],widths=[1.15,3.55,2.05],font_size=8.0)

# Lifecycle and tracking
doc.add_page_break(); add_heading(doc,'4  Planning and tracking lifecycle',1)
add_heading(doc,'From goal to Jira work',2)
add_p(doc,'A planning request includes a goal, date window, optional backlog items, exclusions and constraints. The Team Context Agent first reads the configured Jira board and project data. Intake normalizes the request; subsequent agents use the resulting shared state and evidence. The Orchestrator forms the plan record, and Jira Publisher creates the work.')
add_table(doc,['Stage','What happens','Result'],[
 ('1  Request','Browser submits a goal, dates and optional task list.','A request object with goal and constraints.'),
 ('2  Context','Server validates model, Jira board and confirmed team profiles; Team Context captures Jira and roster evidence.','A timestamped snapshot and fingerprint.'),
 ('3  Plan','LangGraph runs Intake → Decomposer → Planner → Predictor → Assignment → Monitor → Evaluator → Orchestrator.','A selected scope, task owners, estimates, criteria, and summaries.'),
 ('4  Publish','Jira Publisher creates tasks/subtasks, transitions each issue to To Do when available, assigns mapped accounts, adds sprint membership when configured and links dependencies.','Jira issue keys and per-operation success/failure.'),
 ('5  Track','User selects Track on a saved plan; server retrieves plan-labeled Jira issues and runs the tracking graph.','Monitor report, status table, evidence, completion bar and stored tracking record.')
],widths=[1.08,3.55,2.12],font_size=8.2)
add_heading(doc,'Tracking is separate from planning',2)
add_figure(doc,tracking,'Figure 4. Track refreshes progress for a saved plan; it does not repeat the full planning graph.',6.88)
add_p(doc,'The tracking graph runs Monitor Agent and then an Orchestrator summary. In the UI, the tracking panel presents the Monitor Agent output and task evidence. Completion is counted from Jira’s Done status category or a resolution date. KPI rows appear only when tracking metrics are configured.')

# API and data
doc.add_page_break(); add_heading(doc,'5  APIs, data and reporting',1)
add_heading(doc,'Primary application routes',2)
add_table(doc,['Route','Purpose'],[
 ('GET /api/status','Reports readiness for model, Jira board and team profile setup.'),
 ('GET /api/team · PUT /api/team/:id','Loads project team context and saves profile edits.'),
 ('GET /api/settings · PUT /api/settings','Reads and updates workspace, Jira project/board/sprint and model options.'),
 ('GET /api/plans','Returns saved plans and their latest tracking report.'),
 ('POST /api/plans/stream','Starts planning and streams agent progress events as Server-Sent Events.'),
 ('POST /api/plans/:id/track','Checks Jira progress for a saved plan and persists Monitor output.'),
 ('PUT /api/plans/:id','Saves plan task edits.'),
 ('DELETE /api/plans/:id','Removes a plan from NovaSprint history; Jira issues are not deleted.'),
 ('GET /api/audit','Returns recent workspace audit events.')
],widths=[2.3,4.45],font_size=8.3)
add_heading(doc,'Shared planning state',2)
add_p(doc,'The LangGraph state holds the request and workspace context, Jira snapshot, normalized brief, work items, scope candidates, predictions, assignments, findings, evaluation, final decision, plan, agent runs and agent outputs. Reducers append run records and outputs as graph nodes complete. A plan snapshot and team fingerprint support freshness checks before later publishing operations.')
add_heading(doc,'Plan export',2)
add_p(doc,'The current Download Task List action creates a standalone HTML file. It includes the sprint goal, dates, task names, owners, descriptions, acceptance criteria and subtasks. This focused export is intended to be easy to save, print or share without adding the full agent-review page.')
add_callout(doc,'Data storage','The checked-in seed file initializes a new local workspace. Runtime settings, plans, tracking reports and audit events are saved to data/workspace.json; that live workspace file and .env credentials are ignored by Git.',GOLD)

# Model, safeguards and ops
doc.add_page_break(); add_heading(doc,'6  Models, safeguards and operating boundaries',1)
add_heading(doc,'Model configuration',2)
add_p(doc,'lib/model.js selects the provider from MODEL_PROVIDER. Nebius Token Factory is configured through its OpenAI-compatible API using ChatOpenAI, with model name, base URL and key supplied through environment configuration. Ollama is an alternative local provider using ChatOllama. The model is set to temperature zero and uses one retry for provider requests; the role runner has its own compact retry for structured-output length failures.')
add_heading(doc,'Evidence and failure handling',2)
add_bullets(doc,[
 'Team Context requires Jira credentials, a project and a selected board; planning connection checks also require a reachable model and confirmed team profiles with skills.',
 'Agent responses are parsed against Zod schemas. When output is truncated, the runner retries once with compact context and an increased response budget; a second failure becomes an explicit agent error.',
 'Predictor estimates depend on evidence quality. When capacity or velocity baselines are missing, the plan identifies that limitation instead of treating estimates as a committed forecast.',
 'Evaluator does not invent results. It marks usability checks as not evidenced when no completed implementation or user-test evidence was supplied, and only compares metrics that are present in source evidence.',
 'Jira issue creation is followed by a transition to To Do. This depends on the project workflow exposing a transition to a status named To Do; if not, the publisher records the failure and returns a partial result.',
 'Track reports what Jira currently shows. It cannot determine whether a product KPI improved without configured metrics and post-change measurement data.'
])
add_heading(doc,'Security and deployment notes',2)
add_p(doc,'The server binds to loopback (127.0.0.1) by default and loads secrets from the server-side .env file. Do not commit .env or data/workspace.json. This implementation stores a single workspace in JSON rather than a multi-user database; it is designed as a local team tool, not a hosted multi-tenant service with user authentication.')
add_table(doc,['Operational topic','Current behavior'],[
 ('Provider secrets','Read on the server; the browser receives provider readiness metadata, not API keys.'),
 ('Local persistence','Atomic temp-file write followed by rename; audit history is capped at 500 events.'),
 ('Jira task state','Publisher targets To Do via Jira workflow transitions; actual transition options are Jira project configuration.'),
 ('Plan deletion','Deletes the saved NovaSprint plan record only. Existing Jira issues are preserved.'),
 ('Outcome metrics','Requires configured KPI definitions and measured evidence; absence is represented as unavailable.')
],widths=[1.65,5.1],font_size=8.4)

add_heading(doc,'Implementation summary',1)
add_p(doc,'NovaSprint’s key architectural choice is to make planning a coordinated sequence of specialized, evidence-bounded agent roles. LangGraph supplies the control flow and shared state; LangChain adapters supply configurable model access; Zod schemas enforce structured outputs; the Node.js service gates external access and persists plans; Jira remains the source of truth for delivery status. The separation lets users inspect why work was proposed, publish it, and later track observed progress.')
add_heading(doc,'Code map',2)
add_table(doc,['File','Responsibility'],[
 ('lib/graph.js','Agent prompts, runRole, planning graph, tracking graph, publishing and output assembly.'),
 ('lib/agent-schemas.js','Structured output schemas and response-size contracts.'),
 ('lib/model.js · lib/jira.js','Model provider configuration and Jira REST/Agile client functions.'),
 ('server.js · lib/store.js','HTTP API, SSE stream, request routing, JSON persistence and audit log.'),
 ('app.js · planner-ui.js · index.html · styles.css','Team and history screens, planner interactions, task-list export, responsive layout and tracking display.')
],widths=[2.65,4.1],font_size=8.2)

doc.core_properties.title='NovaSprint Implementation Guide'
doc.core_properties.subject='Brain-inspired MAP planning architecture and technical implementation'
doc.core_properties.author='NovaSprint'
doc.core_properties.keywords='NovaSprint, MAP, LangGraph, sprint planning, Jira, technical architecture'
doc.save(OUT)
shutil.rmtree(TMP,ignore_errors=True)
print(OUT)
